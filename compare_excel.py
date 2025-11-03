import sys
from collections import OrderedDict
from pathlib import Path

try:
    from tkinter import Tk, messagebox
    from tkinter import filedialog
except ImportError as exc:  # pragma: no cover - tkinter might not be available
    raise SystemExit("tkinter is required to run this script") from exc

from openpyxl import load_workbook
from tqdm import tqdm


def select_file(title: str) -> Path:
    """Open a file dialog to select an Excel file."""

    root = Tk()
    root.withdraw()
    file_path = filedialog.askopenfilename(
        title=title,
        filetypes=[("Excel files", "*.xlsx *.xlsm *.xltx *.xltm"), ("All files", "*.*")],
    )
    root.destroy()

    if not file_path:
        raise SystemExit("No file selected")

    return Path(file_path)


def ask_save_path(default_name: str) -> Path:
    """Ask the user for a destination to save the result workbook."""

    root = Tk()
    root.withdraw()
    file_path = filedialog.asksaveasfilename(
        initialfile=default_name,
        defaultextension=".xlsx",
        filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
    )
    root.destroy()

    if not file_path:
        raise SystemExit("No save path selected")

    return Path(file_path)


def build_header_map(sheet) -> OrderedDict:
    """Return an ordered mapping of header name to column index for the first row."""

    headers = OrderedDict()
    for cell in sheet[1]:
        headers[cell.value] = cell.column
    return headers


def collect_sheet_data(sheet, tag_col, korean_col, chinese_col):
    """Collect mapping from tag to (korean, chinese) from a sheet."""

    data = {}
    max_row = sheet.max_row
    for row_idx in range(2, max_row + 1):
        tag_value = sheet.cell(row=row_idx, column=tag_col).value
        if tag_value in data:
            # Latest entry wins; duplicates are overwritten intentionally.
            pass
        korean_value = sheet.cell(row=row_idx, column=korean_col).value
        chinese_value = sheet.cell(row=row_idx, column=chinese_col).value
        data[tag_value] = (korean_value, chinese_value)
    return data


def compare_sheets(sheet_a, sheet_b, progress_bar=None):
    """Compare corresponding sheets and populate old value columns."""

    headers_a = build_header_map(sheet_a)

    required_columns = {"Tag", "Korean", "Chinese"}
    if not required_columns.issubset(headers_a.keys()):
        missing = required_columns - set(headers_a.keys())
        raise ValueError(f"Sheet '{sheet_a.title}' is missing columns: {', '.join(missing)}")

    korean_col = headers_a["Korean"]
    sheet_a.insert_cols(korean_col + 1)
    sheet_a.cell(row=1, column=korean_col + 1).value = "旧Korean"

    headers_a = build_header_map(sheet_a)
    chinese_col = headers_a["Chinese"]
    sheet_a.insert_cols(chinese_col + 1)
    sheet_a.cell(row=1, column=chinese_col + 1).value = "旧Chinese"

    headers_a = build_header_map(sheet_a)
    tag_col = headers_a["Tag"]
    old_korean_col = headers_a["旧Korean"]
    old_chinese_col = headers_a["旧Chinese"]
    korean_col = headers_a["Korean"]
    chinese_col = headers_a["Chinese"]

    headers_b = build_header_map(sheet_b)
    if not required_columns.issubset(headers_b.keys()):
        missing = required_columns - set(headers_b.keys())
        raise ValueError(f"Sheet '{sheet_b.title}' is missing columns: {', '.join(missing)}")

    tag_col_b = headers_b["Tag"]
    korean_col_b = headers_b["Korean"]
    chinese_col_b = headers_b["Chinese"]

    lookup_b = collect_sheet_data(sheet_b, tag_col_b, korean_col_b, chinese_col_b)

    for row_idx in range(2, sheet_a.max_row + 1):
        tag_value = sheet_a.cell(row=row_idx, column=tag_col).value
        korean_value_a = sheet_a.cell(row=row_idx, column=korean_col).value
        chinese_value_a = sheet_a.cell(row=row_idx, column=chinese_col).value

        b_values = lookup_b.get(tag_value)
        if b_values is None:
            sheet_a.cell(row=row_idx, column=old_korean_col).value = "新增内容"
            sheet_a.cell(row=row_idx, column=old_chinese_col).value = "新增内容"
            if progress_bar is not None:
                progress_bar.update(1)
            continue

        korean_value_b, chinese_value_b = b_values

        if korean_value_a == korean_value_b:
            sheet_a.cell(row=row_idx, column=old_korean_col).value = "无变更"
        else:
            sheet_a.cell(row=row_idx, column=old_korean_col).value = korean_value_b

        if chinese_value_a == chinese_value_b:
            sheet_a.cell(row=row_idx, column=old_chinese_col).value = "无变更"
        else:
            sheet_a.cell(row=row_idx, column=old_chinese_col).value = chinese_value_b

        if progress_bar is not None:
            progress_bar.update(1)


def resolve_workbooks():
    """Resolve workbook paths from CLI arguments or the file dialog."""

    args = sys.argv[1:]
    if not args:
        file_a = select_file("选择文档A (Vn+1)")
        file_b = select_file("选择文档B (Vn)")
        return file_a, file_b

    if len(args) != 2:
        raise SystemExit("Usage: python compare_excel.py [file_a file_b]")

    file_a = Path(args[0])
    file_b = Path(args[1])

    if not file_a.exists():
        raise SystemExit(f"文件不存在: {file_a}")
    if not file_b.exists():
        raise SystemExit(f"文件不存在: {file_b}")

    return file_a, file_b


def main():
    file_a, file_b = resolve_workbooks()

    wb_a = load_workbook(filename=file_a)
    wb_b = load_workbook(filename=file_b)

    common_sheets = [name for name in wb_a.sheetnames if name in wb_b.sheetnames]

    total_rows = sum(max(wb_a[sheet].max_row - 1, 0) for sheet in common_sheets)
    progress_bar = tqdm(total=total_rows, desc="对比进度", unit="行")

    try:
        for sheet_name in common_sheets:
            sheet_a = wb_a[sheet_name]
            sheet_b = wb_b[sheet_name]
            progress_bar.set_postfix_str(sheet_name)
            compare_sheets(sheet_a, sheet_b, progress_bar)
    finally:
        progress_bar.close()

    root = Tk()
    root.withdraw()
    messagebox.showinfo("完成", "对比完成")
    root.destroy()

    save_path = ask_save_path("AB文档对比结果.xlsx")
    wb_a.save(save_path)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:  # pragma: no cover - surface errors to the user
        root = Tk()
        root.withdraw()
        messagebox.showerror("错误", str(exc))
        root.destroy()
        sys.exit(1)
